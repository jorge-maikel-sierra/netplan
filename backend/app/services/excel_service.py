"""Excel parsing and export service for node imports and project exports."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from typing import List, Tuple, Dict, Any, Optional
from io import BytesIO


REQUIRED_COLUMNS = ["nombre", "latitud", "longitud"]
EDGE_REQUIRED_COLUMNS = ["origen", "destino", "costo"]
EDGE_VALID_TYPES = {"normal", "mandatory", "forbidden"}
MAX_ROWS = 10_000
VALID_MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def validate_mime_type(mime_type: str) -> None:
    """Validate that the file is an .xlsx file by MIME type."""
    if mime_type != VALID_MIME_TYPE:
        raise ValueError("UNSUPPORTED_FILE_TYPE")


def parse_nodes_xlsx(file_bytes: bytes) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    """
    Parse an .xlsx file containing node data.
    
    Args:
        file_bytes: Raw bytes of the .xlsx file
        
    Returns:
        Tuple of (valid_rows, errors)
        - valid_rows: List of dicts with keys: nombre, latitud, longitud, tipo
        - errors: List of dicts with keys: row, field, message
    """
    valid_rows = []
    errors = []
    
    try:
        workbook = openpyxl.load_workbook(BytesIO(file_bytes), read_only=True)
        sheet = workbook.active
    except Exception as e:
        errors.append({"row": 0, "field": "file", "message": f"Invalid .xlsx file: {str(e)}"})
        return valid_rows, errors
    
    # Get headers from first row
    headers = []
    for cell in sheet[1]:
        headers.append(cell.value)
    
    # Validate required columns
    missing_columns = [col for col in REQUIRED_COLUMNS if col not in headers]
    if missing_columns:
        errors.append({
            "row": 0, 
            "field": "columns", 
            "message": f"Missing required columns: {', '.join(missing_columns)}. Required: {', '.join(REQUIRED_COLUMNS)}"
        })
        return valid_rows, errors
    
    # Map column names to indices
    col_indices = {col: headers.index(col) for col in REQUIRED_COLUMNS}
    tipo_idx = headers.index("tipo") if "tipo" in headers else None
    
    row_num = 1  # 1-indexed for user-friendly error messages (header is row 1)
    
    for row in sheet.iter_rows(min_row=2, values_only=True):
        row_num += 1
        
        # Check row limit
        if row_num > MAX_ROWS + 1:
            errors.append({
                "row": row_num,
                "field": "row_count",
                "message": f"Maximum {MAX_ROWS} rows allowed. File has more rows."
            })
            break
        
        # Skip completely empty rows
        if all(cell is None for cell in row):
            continue
        
        row_errors = []
        
        # Parse required fields
        nombre = row[col_indices["nombre"]]
        latitud = row[col_indices["latitud"]]
        longitud = row[col_indices["longitud"]]
        
        # Validate nombre
        if not nombre or not str(nombre).strip():
            row_errors.append({"row": row_num, "field": "nombre", "message": "Nombre is required"})
        else:
            nombre = str(nombre).strip()
        
        # Validate latitud
        if latitud is None:
            row_errors.append({"row": row_num, "field": "latitud", "message": "Latitud is required"})
        else:
            try:
                lat = float(latitud)
                if lat < -90 or lat > 90:
                    row_errors.append({"row": row_num, "field": "latitud", "message": "Latitud must be between -90 and 90"})
                else:
                    latitud = lat
            except (ValueError, TypeError):
                row_errors.append({"row": row_num, "field": "latitud", "message": "Latitud must be a valid number"})
        
        # Validate longitud
        if longitud is None:
            row_errors.append({"row": row_num, "field": "longitud", "message": "Longitud is required"})
        else:
            try:
                lng = float(longitud)
                if lng < -180 or lng > 180:
                    row_errors.append({"row": row_num, "field": "longitud", "message": "Longitud must be between -180 and 180"})
                else:
                    longitud = lng
            except (ValueError, TypeError):
                row_errors.append({"row": row_num, "field": "longitud", "message": "Longitud must be a valid number"})
        
        # Parse optional tipo
        tipo = "city"
        if tipo_idx is not None and tipo_idx < len(row):
            tipo_val = row[tipo_idx]
            if tipo_val and str(tipo_val).strip():
                tipo_str = str(tipo_val).strip().lower()
                if tipo_str in ("city", "tower", "datacenter"):
                    tipo = tipo_str
                else:
                    row_errors.append({"row": row_num, "field": "tipo", "message": "Tipo must be 'city', 'tower', or 'datacenter'"})
        
        if row_errors:
            errors.extend(row_errors)
        else:
            valid_rows.append({
                "nombre": nombre,
                "latitud": latitud,
                "longitud": longitud,
                "tipo": tipo
            })
    
    return valid_rows, errors


def parse_edges_xlsx(file_bytes: bytes) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    """
    Parse an .xlsx file containing edge data.

    The function looks for a second sheet first; if not found or lacking required
    columns, it falls back to the first sheet.  Required columns: origen, destino,
    costo.  Optional column: tipo_restriccion (defaults to 'normal').

    Args:
        file_bytes: Raw bytes of the .xlsx file

    Returns:
        Tuple of (valid_rows, errors)
        - valid_rows: List of dicts with keys: origen, destino, costo, tipo_restriccion
        - errors: List of dicts with keys: row, field, message
    """
    valid_rows: List[Dict[str, Any]] = []
    errors: List[Dict[str, Any]] = []

    # Load workbook
    try:
        workbook = openpyxl.load_workbook(BytesIO(file_bytes), read_only=True)
    except Exception as e:
        errors.append({"row": 0, "field": "file", "message": f"Invalid .xlsx file: {e}"})
        return valid_rows, errors

    # Choose the sheet to parse edges from
    sheet = None
    if len(workbook.sheetnames) >= 2:
        # Try the second sheet
        candidate = workbook[workbook.sheetnames[1]]
        headers = [cell.value for cell in next(candidate.iter_rows(min_row=1, max_row=1))]
        if all(col in headers for col in EDGE_REQUIRED_COLUMNS):
            sheet = candidate

    if sheet is None:
        # Fallback to the active (first) sheet
        sheet = workbook.active
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]

    # Validate required columns
    missing = [col for col in EDGE_REQUIRED_COLUMNS if col not in headers]
    if missing:
        errors.append({
            "row": 0,
            "field": "columns",
            "message": f"Missing required edge columns: {', '.join(missing)}. "
                       f"Required: {', '.join(EDGE_REQUIRED_COLUMNS)}",
        })
        return valid_rows, errors

    # Map column names to indices
    col_idx = {col: headers.index(col) for col in EDGE_REQUIRED_COLUMNS}
    tipo_idx = headers.index("tipo_restriccion") if "tipo_restriccion" in headers else None

    row_num = 1  # header row
    for row in sheet.iter_rows(min_row=2, values_only=True):
        row_num += 1

        if row_num > MAX_ROWS + 1:
            errors.append({
                "row": row_num,
                "field": "row_count",
                "message": f"Maximum {MAX_ROWS} rows allowed. File has more rows.",
            })
            break

        # Skip completely empty rows
        if all(cell is None for cell in row):
            continue

        row_errors: List[Dict[str, Any]] = []

        # ---- origen ----
        origen = row[col_idx["origen"]]
        if origen is None or not str(origen).strip():
            row_errors.append({"row": row_num, "field": "origen", "message": "Origen is required"})
        else:
            origen = str(origen).strip()

        # ---- destino ----
        destino = row[col_idx["destino"]]
        if destino is None or not str(destino).strip():
            row_errors.append({"row": row_num, "field": "destino", "message": "Destino is required"})
        else:
            destino = str(destino).strip()

        # ---- costo ----
        costo = row[col_idx["costo"]]
        if costo is None:
            row_errors.append({"row": row_num, "field": "costo", "message": "Costo is required"})
        else:
            try:
                costo_val = float(costo)
                if costo_val < 0:
                    row_errors.append({"row": row_num, "field": "costo", "message": "Costo must be non-negative"})
                else:
                    costo = costo_val
            except (ValueError, TypeError):
                row_errors.append({"row": row_num, "field": "costo", "message": "Costo must be a valid number"})

        # ---- tipo_restriccion (optional) ----
        tipo = "normal"
        if tipo_idx is not None and tipo_idx < len(row):
            tipo_val = row[tipo_idx]
            if tipo_val and str(tipo_val).strip():
                tipo_str = str(tipo_val).strip().lower()
                if tipo_str in EDGE_VALID_TYPES:
                    tipo = tipo_str
                else:
                    row_errors.append({
                        "row": row_num,
                        "field": "tipo_restriccion",
                        "message": f"Tipo must be one of: {', '.join(sorted(EDGE_VALID_TYPES))}",
                    })

        if row_errors:
            errors.extend(row_errors)
        else:
            valid_rows.append({
                "row_num": row_num,
                "origen": origen,
                "destino": destino,
                "costo": costo,
                "tipo_restriccion": tipo,
            })

    return valid_rows, errors


# ---------------------------------------------------------------------------
# Export workbook
# ---------------------------------------------------------------------------

HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center")


def _style_header(ws: openpyxl.worksheet.worksheet.Worksheet) -> None:
    """Apply bold white-on-blue styling to the header row and freeze it."""
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
    ws.freeze_panes = "A2"


def _auto_width(ws: openpyxl.worksheet.worksheet.Worksheet) -> None:
    """Set column widths based on max content length (capped at 50)."""
    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            val = str(cell.value) if cell.value is not None else ""
            max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = min(max_len + 3, 50)


def _build_edge_id_set(mst_result: Optional[dict]) -> set[str]:
    """Extract the set of edge IDs present in the MST result."""
    if mst_result is None:
        return set()
    raw = mst_result.get("edge_ids") or mst_result.get("mst_edges")
    if isinstance(raw, list):
        return {str(eid) for eid in raw}
    return set()


def build_export_workbook(
    project: dict,
    nodes: list[dict],
    edges: list[dict],
    mst_result: Optional[dict],
) -> BytesIO:
    """Build a styled 3-sheet .xlsx workbook for project export.

    Sheets:
      - Project Info: name, description, node/edge counts, MST total cost
      - Nodes: label, type, latitude, longitude (sorted by label)
      - Edges & MST: node_a, node_b, cost, constraint, in-MST flag

    Returns:
        BytesIO containing the .xlsx data (ready for StreamingResponse).
    """
    wb = openpyxl.Workbook()

    # -------------------------------------------------------------------
    # Sheet 1: Project Info
    # -------------------------------------------------------------------
    ws_info = wb.active
    ws_info.title = "Project Info"
    info_fields = [
        ("Project Name", project.get("name", "")),
        ("Description", project.get("description", "")),
        ("Nodes", str(len(nodes))),
        ("Edges", str(len(edges))),
    ]
    if mst_result:
        total_cost = mst_result.get("total_cost", 0)
        info_fields.append(("Total Cost", str(total_cost)))

    for i, (key, val) in enumerate(info_fields, start=1):
        ws_info.cell(row=i, column=1, value=f"{key}:")
        ws_info.cell(row=i, column=2, value=val)

    _style_header(ws_info)
    _auto_width(ws_info)

    # -------------------------------------------------------------------
    # Sheet 2: Nodes
    # -------------------------------------------------------------------
    ws_nodes = wb.create_sheet("Nodes")
    node_headers = ["Label", "Type", "Latitude", "Longitude"]
    ws_nodes.append(node_headers)

    sorted_nodes = sorted(nodes, key=lambda n: str(n.get("name", "")))
    for n in sorted_nodes:
        ws_nodes.append([
            n.get("name", ""),
            n.get("type", ""),
            n.get("lat", ""),
            n.get("lng", ""),
        ])

    _style_header(ws_nodes)
    _auto_width(ws_nodes)

    # -------------------------------------------------------------------
    # Sheet 3: Edges & MST
    # -------------------------------------------------------------------
    ws_edges = wb.create_sheet("Edges & MST")
    edge_headers = ["Node A", "Node B", "Cost", "Constraint", "In MST"]
    ws_edges.append(edge_headers)

    mst_edge_ids = _build_edge_id_set(mst_result)
    # Build a lookup: node_id -> node name
    node_name_by_id: dict[str, str] = {
        str(n.get("id", "")): str(n.get("name", "")) for n in nodes
    }

    sorted_edges = sorted(edges, key=lambda e: float(e.get("cost", 0)), reverse=True)
    for e in sorted_edges:
        node_a_id = str(e.get("node_a_id", ""))
        node_b_id = str(e.get("node_b_id", ""))
        edge_id = str(e.get("id", ""))
        in_mst = "Yes" if edge_id in mst_edge_ids else "No"
        ws_edges.append([
            node_name_by_id.get(node_a_id, node_a_id[:8]),
            node_name_by_id.get(node_b_id, node_b_id[:8]),
            e.get("cost", ""),
            e.get("constraint_type", "normal"),
            in_mst,
        ])

    _style_header(ws_edges)
    _auto_width(ws_edges)

    # Return as BytesIO
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf